import os
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from flask import Flask, render_template, request, redirect, url_for, session, send_file, flash
from models import db, User, Class, Subject, Chapter, QuestionSet, Question, TestResult, StudentAnswer

app = Flask(__name__)
app.secret_key = 'super_secret_exam_key'

# ----------------- DATABASE CONFIGURATION -----------------
# Connects to Render PostgreSQL in production, falls back to SQLite locally
raw_db_url = os.environ.get('DATABASE_URL', 'sqlite:///database.db')

if raw_db_url.startswith("postgres://"):
    raw_db_url = raw_db_url.replace("postgres://", "postgresql+psycopg://", 1)
elif raw_db_url.startswith("postgresql://"):
    raw_db_url = raw_db_url.replace("postgresql://", "postgresql+psycopg://", 1)

app.config['SQLALCHEMY_DATABASE_URI'] = raw_db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db.init_app(app)

# ----------------- DATABASE AUTO-INITIALIZATION & SEEDING -----------------

with app.app_context():
    db.create_all()

    # Pre-populate all standard and stream-specific classes
    all_classes = [
        "Class 6", "Class 7", "Class 8", "Class 9", "Class 10",
        "11th Science", "11th Commerce", "11th Arts",
        "12th Science", "12th Commerce", "12th Arts"
    ]
    for c_name in all_classes:
        if not Class.query.filter_by(class_name=c_name).first():
            db.session.add(Class(class_name=c_name))
    db.session.commit()

    # Pre-populate default teacher account
    if not User.query.filter_by(email="teacher@test.com").first():
        default_teacher = User(
            full_name="Prof. Sharma",
            email="teacher@test.com",
            password="admin",
            role="teacher"
        )
        db.session.add(default_teacher)
        db.session.commit()

    # Pre-populate default sample student
    class_10 = Class.query.filter_by(class_name="Class 10").first()
    if class_10 and not User.query.filter_by(email="rahul@test.com").first():
        default_student = User(
            full_name="Rahul Verma",
            email="rahul@test.com",
            password="123",
            role="student",
            roll_number="101",
            class_id=class_10.id
        )
        db.session.add(default_student)
        db.session.commit()

    # Pre-populate sample curriculum if empty
    if class_10 and not Subject.query.filter_by(class_id=class_10.id, subject_name="Mathematics").first():
        math = Subject(class_id=class_10.id, subject_name="Mathematics")
        db.session.add(math)
        db.session.commit()

        chap = Chapter(subject_id=math.id, chapter_no=1, title="Linear Equations")
        db.session.add(chap)
        db.session.commit()

        q_set = QuestionSet(chapter_id=chap.id, set_title="Set A", time_limit=10, total_marks=10)
        db.session.add(q_set)
        db.session.commit()

        sample_questions = [
            Question(
                set_id=q_set.id,
                question_text=f"Question {i}: What is the value of x if x + {i} = {i * 2}?",
                option_a=str(i),
                option_b=str(i + 1),
                option_c=str(i + 2),
                option_d=str(i + 3),
                correct_option="A"
            )
            for i in range(1, 11)
        ]
        db.session.add_all(sample_questions)
        db.session.commit()

# ----------------- AUTHENTICATION & REGISTRATION -----------------

@app.route('/')
def index():
    if 'user_id' in session:
        return redirect(url_for('teacher_dashboard' if session.get('role') == 'teacher' else 'select_exam'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email = request.form.get('email')
        password = request.form.get('password')
        
        user = User.query.filter_by(email=email, password=password).first()
        if user:
            session['user_id'] = user.id
            session['full_name'] = user.full_name
            session['role'] = user.role
            
            if user.role == 'teacher':
                return redirect(url_for('teacher_dashboard'))
            return redirect(url_for('select_exam'))
        else:
            flash('Invalid email or password', 'danger')
            
    return render_template('login.html')

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        full_name = request.form.get('full_name')
        roll_number = request.form.get('roll_number')
        email = request.form.get('email')
        password = request.form.get('password')
        class_id = request.form.get('class_id', type=int)

        if User.query.filter_by(email=email).first():
            flash('Email already registered. Please log in.', 'danger')
            return redirect(url_for('register'))
        
        if User.query.filter_by(roll_number=roll_number).first():
            flash('Roll number already exists. Check your details.', 'danger')
            return redirect(url_for('register'))

        new_student = User(
            full_name=full_name,
            email=email,
            password=password,
            role='student',
            roll_number=roll_number,
            class_id=class_id
        )
        db.session.add(new_student)
        db.session.commit()

        flash('Registration successful! Please log in.', 'success')
        return redirect(url_for('login'))

    classes = Class.query.order_by(Class.id).all()
    return render_template('register.html', classes=classes)

@app.route('/register-teacher', methods=['GET', 'POST'])
def register_teacher():
    SECRET_KEY_PHRASE = "TEACHER2026"

    if request.method == 'POST':
        full_name = request.form.get('full_name')
        email = request.form.get('email')
        password = request.form.get('password')
        secret_code = request.form.get('secret_code')

        if secret_code != SECRET_KEY_PHRASE:
            flash("Invalid teacher authorization code. Access denied.", "danger")
            return redirect(url_for('register_teacher'))

        if User.query.filter_by(email=email).first():
            flash("Email already registered.", "danger")
            return redirect(url_for('register_teacher'))

        new_teacher = User(
            full_name=full_name,
            email=email,
            password=password,
            role='teacher',
            roll_number=None,
            class_id=None
        )
        db.session.add(new_teacher)
        db.session.commit()

        flash("Teacher account created successfully! Please log in.", "success")
        return redirect(url_for('login'))

    return render_template('register_teacher.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

# ----------------- STUDENT FLOW -----------------

@app.route('/select-exam')
def select_exam():
    if 'user_id' not in session or session.get('role') != 'student':
        return redirect(url_for('login'))
    
    current_uid = int(session['user_id'])
    student = User.query.get(current_uid)
    student_class = Class.query.get(student.class_id) if student.class_id else None
    subjects = Subject.query.filter_by(class_id=student.class_id).all() if student.class_id else []

    selected_subject_id = request.args.get('subject_id', type=int)
    selected_chapter_id = request.args.get('chapter_id', type=int)

    chapters = Chapter.query.filter_by(subject_id=selected_subject_id).all() if selected_subject_id else []
    sets = QuestionSet.query.filter_by(chapter_id=selected_chapter_id).all() if selected_chapter_id else []

    attempts = TestResult.query.filter_by(student_id=current_uid).all()
    attempted_set_ids = [a.set_id for a in attempts]

    return render_template(
        'select_exam.html',
        student_class=student_class,
        subjects=subjects,
        chapters=chapters,
        sets=sets,
        selected_subject_id=selected_subject_id,
        selected_chapter_id=selected_chapter_id,
        attempted_set_ids=attempted_set_ids
    )

@app.route('/exam/<int:set_id>', methods=['GET', 'POST'])
def exam(set_id):
    if 'user_id' not in session or session.get('role') != 'student':
        return redirect(url_for('login'))

    current_uid = int(session['user_id'])
    student = User.query.get(current_uid)
    q_set = QuestionSet.query.get_or_404(set_id)

    existing_attempt = TestResult.query.filter_by(
        student_id=current_uid,
        set_id=q_set.id
    ).first()

    if existing_attempt:
        flash("Test already attempted! You can only attempt a test once.", "warning")
        return redirect(url_for('result', result_id=existing_attempt.id))

    chapter = Chapter.query.get(q_set.chapter_id)
    subject = Subject.query.get(chapter.subject_id)
    if student.class_id and subject.class_id != student.class_id:
        flash("Unauthorized: You cannot access assessments outside your assigned class.", "danger")
        return redirect(url_for('select_exam'))

    questions = Question.query.filter_by(set_id=set_id).all()

    if request.method == 'POST':
        score = 0
        time_taken = request.form.get('time_taken', default=0, type=int)

        result = TestResult(
            student_id=current_uid,
            set_id=q_set.id,
            score=0,
            time_taken_seconds=time_taken
        )
        db.session.add(result)
        db.session.commit()

        answers_to_store = []
        for q in questions:
            user_choice = request.form.get(f'question_{q.id}')
            is_correct = (user_choice == q.correct_option) if user_choice else False
            
            if is_correct:
                score += 1

            answers_to_store.append(
                StudentAnswer(
                    result_id=result.id,
                    question_id=q.id,
                    selected_option=user_choice,
                    is_correct=is_correct
                )
            )

        result.score = score
        db.session.add_all(answers_to_store)
        db.session.commit()
        
        return redirect(url_for('result', result_id=result.id))

    return render_template('exam.html', q_set=q_set, questions=questions)

@app.route('/result/<int:result_id>')
def result(result_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
        
    res = TestResult.query.get_or_404(result_id)
    answers = StudentAnswer.query.filter_by(result_id=res.id).all()

    top_performers = TestResult.query.filter_by(set_id=res.set_id)\
        .order_by(TestResult.score.desc(), TestResult.time_taken_seconds.asc())\
        .limit(5)\
        .all()
    
    return render_template('result.html', result=res, answers=answers, top_performers=top_performers)

# ----------------- TEACHER ACTIONS & QUESTION BANK -----------------

@app.route('/teacher/dashboard')
def teacher_dashboard():
    if 'user_id' not in session or session.get('role') != 'teacher':
        return redirect(url_for('login'))

    classes = Class.query.order_by(Class.id).all()
    selected_class_id = request.args.get('class_id', type=int)
    students = User.query.filter_by(role='student', class_id=selected_class_id).all() if selected_class_id else []
    selected_student_id = request.args.get('student_id', type=int)
    
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')

    query = TestResult.query

    if selected_student_id:
        query = query.filter_by(student_id=selected_student_id)
    elif selected_class_id:
        student_ids = [s.id for s in students]
        query = query.filter(TestResult.student_id.in_(student_ids))

    if start_date_str:
        try:
            start_dt = datetime.strptime(start_date_str, "%Y-%m-%d")
            query = query.filter(TestResult.submitted_at >= start_dt)
        except ValueError:
            pass

    if end_date_str:
        try:
            end_dt = datetime.strptime(end_date_str + " 23:59:59", "%Y-%m-%d %H:%M:%S")
            query = query.filter(TestResult.submitted_at <= end_dt)
        except ValueError:
            pass

    results = query.order_by(TestResult.submitted_at.desc()).all()

    student_summary = None
    if selected_student_id and results:
        total_tests = len(results)
        total_scored = sum(r.score for r in results)
        total_possible = total_tests * 10
        percentage = round((total_scored / total_possible) * 100, 2) if total_possible > 0 else 0
        student_summary = {
            "total_tests": total_tests,
            "total_scored": total_scored,
            "total_possible": total_possible,
            "percentage": percentage
        }

    return render_template(
        'teacher_dashboard.html',
        classes=classes,
        students=students,
        results=results,
        selected_class_id=selected_class_id,
        selected_student_id=selected_student_id,
        start_date=start_date_str,
        end_date=end_date_str,
        student_summary=student_summary
    )

@app.route('/teacher/manage-questions')
def manage_questions():
    if 'user_id' not in session or session.get('role') != 'teacher':
        return redirect(url_for('login'))

    classes = Class.query.order_by(Class.id).all()
    subjects = Subject.query.join(Class).order_by(Class.id, Subject.subject_name).all()
    chapters = Chapter.query.join(Subject).order_by(Subject.id, Chapter.chapter_no).all()

    return render_template('manage_questions.html', classes=classes, subjects=subjects, chapters=chapters)

@app.route('/teacher/add-subject', methods=['POST'])
def add_subject():
    if 'user_id' not in session or session.get('role') != 'teacher':
        return redirect(url_for('login'))

    class_id = request.form.get('class_id', type=int)
    subject_name = request.form.get('subject_name', '').strip()

    if class_id and subject_name:
        new_sub = Subject(class_id=class_id, subject_name=subject_name)
        db.session.add(new_sub)
        db.session.commit()
        flash(f"Subject '{subject_name}' added successfully!", "success")
    else:
        flash("Invalid subject details.", "danger")

    return redirect(url_for('manage_questions'))

@app.route('/teacher/add-chapter', methods=['POST'])
def add_chapter():
    if 'user_id' not in session or session.get('role') != 'teacher':
        return redirect(url_for('login'))

    subject_id = request.form.get('subject_id', type=int)
    chapter_no = request.form.get('chapter_no', type=int)
    title = request.form.get('title', '').strip()

    if subject_id and chapter_no and title:
        new_ch = Chapter(subject_id=subject_id, chapter_no=chapter_no, title=title)
        db.session.add(new_ch)
        db.session.commit()
        flash(f"Chapter '{title}' added successfully!", "success")
    else:
        flash("Invalid chapter details.", "danger")

    return redirect(url_for('manage_questions'))

@app.route('/teacher/add-question-set', methods=['POST'])
def add_question_set():
    if 'user_id' not in session or session.get('role') != 'teacher':
        return redirect(url_for('login'))

    chapter_id = request.form.get('chapter_id', type=int)
    set_title = request.form.get('set_title', '').strip()
    time_limit = request.form.get('time_limit', default=10, type=int)

    if not (chapter_id and set_title):
        flash("Please provide all set details.", "danger")
        return redirect(url_for('manage_questions'))

    q_set = QuestionSet(chapter_id=chapter_id, set_title=set_title, time_limit=time_limit, total_marks=10)
    db.session.add(q_set)
    db.session.commit()

    questions = []
    for i in range(1, 11):
        q_text = request.form.get(f'q_text_{i}')
        optA = request.form.get(f'q_optA_{i}')
        optB = request.form.get(f'q_optB_{i}')
        optC = request.form.get(f'q_optC_{i}')
        optD = request.form.get(f'q_optD_{i}')
        correct = request.form.get(f'q_correct_{i}')

        question = Question(
            set_id=q_set.id,
            question_text=q_text,
            option_a=optA,
            option_b=optB,
            option_c=optC,
            option_d=optD,
            correct_option=correct
        )
        questions.append(question)

    db.session.add_all(questions)
    db.session.commit()

    flash(f"Question Set '{set_title}' with 10 MCQs saved successfully!", "success")
    return redirect(url_for('manage_questions'))

@app.route('/teacher/edit-question-set/<int:set_id>', methods=['GET', 'POST'])
def edit_question_set(set_id):
    if 'user_id' not in session or session.get('role') != 'teacher':
        return redirect(url_for('login'))

    q_set = QuestionSet.query.get_or_404(set_id)
    questions = Question.query.filter_by(set_id=q_set.id).all()
    chapters = Chapter.query.join(Subject).order_by(Subject.id, Chapter.chapter_no).all()

    if request.method == 'POST':
        q_set.chapter_id = request.form.get('chapter_id', type=int)
        q_set.set_title = request.form.get('set_title', '').strip()
        q_set.time_limit = request.form.get('time_limit', default=10, type=int)

        for idx, q in enumerate(questions, start=1):
            q.question_text = request.form.get(f'q_text_{idx}')
            q.option_a = request.form.get(f'q_optA_{idx}')
            q.option_b = request.form.get(f'q_optB_{idx}')
            q.option_c = request.form.get(f'q_optC_{idx}')
            q.option_d = request.form.get(f'q_optD_{idx}')
            q.correct_option = request.form.get(f'q_correct_{idx}')

        db.session.commit()
        flash(f"Question Set '{q_set.set_title}' updated successfully!", "success")
        return redirect(url_for('manage_questions'))

    return render_template('edit_question_set.html', q_set=q_set, questions=questions, chapters=chapters)

@app.route('/teacher/export-excel')
def export_excel():
    if 'user_id' not in session or session.get('role') != 'teacher':
        return redirect(url_for('login'))

    selected_class_id = request.args.get('class_id', type=int)
    selected_student_id = request.args.get('student_id', type=int)
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')

    query = TestResult.query

    if selected_student_id:
        query = query.filter_by(student_id=selected_student_id)
    elif selected_class_id:
        students = User.query.filter_by(role='student', class_id=selected_class_id).all()
        student_ids = [s.id for s in students]
        query = query.filter(TestResult.student_id.in_(student_ids))

    if start_date_str:
        try:
            start_dt = datetime.strptime(start_date_str, "%Y-%m-%d")
            query = query.filter(TestResult.submitted_at >= start_dt)
        except ValueError:
            pass

    if end_date_str:
        try:
            end_dt = datetime.strptime(end_date_str + " 23:59:59", "%Y-%m-%d %H:%M:%S")
            query = query.filter(TestResult.submitted_at <= end_dt)
        except ValueError:
            pass

    results = query.order_by(TestResult.submitted_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Exam Performance"

    bold_font = Font(name="Calibri", size=11, bold=True)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    summary_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    if selected_student_id:
        student = User.query.get(selected_student_id)
        cls = Class.query.get(student.class_id) if (student and student.class_id) else None
        
        total_tests = len(results)
        total_scored = sum(r.score for r in results)
        total_possible = total_tests * 10
        percentage = round((total_scored / total_possible) * 100, 2) if total_possible > 0 else 0

        ws.append(["STUDENT PERFORMANCE REPORT", "", "", ""])
        ws.merge_cells("A1:D1")
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        ws["A1"].fill = header_fill
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.append(["Student Name:", student.full_name if student else "N/A", "Roll Number:", student.roll_number if student else "N/A"])
        ws.append(["Class / Stream:", cls.class_name if cls else "N/A", "Report Date:", datetime.utcnow().strftime("%Y-%m-%d")])
        ws.append(["Date Range:", f"{start_date_str or 'Earliest'} to {end_date_str or 'Latest'}", "", ""])
        ws.append([])

        summary_rows = [
            ["Total Tests Attempted", total_tests],
            ["Total Marks Possible", total_possible],
            ["Total Marks Scored", total_scored],
            ["Overall Percentage", f"{percentage}%"]
        ]
        for row in summary_rows:
            ws.append(row)
            cur_row = ws.max_row
            ws[f"A{cur_row}"].font = bold_font
            ws[f"A{cur_row}"].fill = summary_fill
            ws[f"B{cur_row}"].font = bold_font

        ws.append([])

    headers = [
        "Roll No", "Student Name", "Class", "Subject", 
        "Chapter", "Set", "Score (/10)", "Time Taken (s)", "Submitted At"
    ]
    ws.append(headers)
    table_header_row = ws.max_row

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=table_header_row, column=col_idx)
        cell.font = white_font
        cell.fill = header_fill

    for r in results:
        student = User.query.get(r.student_id)
        q_set = QuestionSet.query.get(r.set_id)
        chap = Chapter.query.get(q_set.chapter_id) if q_set else None
        sub = Subject.query.get(chap.subject_id) if chap else None
        cls = Class.query.get(sub.class_id) if sub else None

        ws.append([
            student.roll_number if student else "N/A",
            student.full_name if student else "Unknown",
            cls.class_name if cls else "N/A",
            sub.subject_name if sub else "N/A",
            chap.title if chap else "N/A",
            q_set.set_title if q_set else "N/A",
            r.score,
            r.time_taken_seconds,
            r.submitted_at.strftime("%Y-%m-%d %H:%M:%S")
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Student_Report_{selected_student_id}.xlsx" if selected_student_id else "Exam_Marks_Filtered.xlsx"

    return send_file(
        stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

if __name__ == '__main__':
    app.run(debug=True)